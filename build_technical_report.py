"""
build_technical_report.py
Generates the comprehensive, publication-grade Technical Report DOCX document for
Traffic Accident Analysis & Severity Risk Modeling: An Exposure-Adjusted Visual Analytics Framework.
Structured and formatted precisely like Impact_of_Generative_AI_Technical_Report_With_Tables_and_Figures.docx.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(BASE_DIR, "report_figures")
OUTPUT_DOCX = os.path.join(BASE_DIR, "Traffic_Accident_Analysis_and_Severity_Risk_Modeling_Technical_Report.docx")

# Styling constants
FONT_NAME = "Times New Roman"
COLOR_TITLE = RGBColor(0x1F, 0x4E, 0x79)   # Navy Blue
COLOR_H1 = RGBColor(0x1F, 0x4E, 0x79)      # Navy Blue
COLOR_H2 = RGBColor(0x2E, 0x74, 0xB5)      # Royal Blue
COLOR_H3 = RGBColor(0x46, 0x46, 0x46)      # Charcoal Dark
COLOR_TEXT = RGBColor(0x28, 0x28, 0x28)    # Body Charcoal
COLOR_MUTED = RGBColor(0x71, 0x80, 0x96)   # Muted Slate
COLOR_BORDER = "CBD5E0"
COLOR_HEADER_BG = "2B6CB0"
COLOR_ROW_ALT = "F7FAFC"
COLOR_ROW_WHITE = "FFFFFF"

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_custom_table(table, col_widths=None, right_align_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=COLOR_BORDER, sz="6")
    
    # Header row
    for c_idx, cell in enumerate(table.rows[0].cells):
        set_cell_background(cell, COLOR_HEADER_BG)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    # Data rows
    right_cols = set(right_align_cols or [])
    for r_idx, row in enumerate(table.rows[1:]):
        bg = COLOR_ROW_WHITE if r_idx % 2 == 0 else COLOR_ROW_ALT
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                if c_idx in right_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = FONT_NAME
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = COLOR_TEXT

    # Apply column widths if provided
    if col_widths:
        for row in table.rows:
            for c_idx, width in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = Inches(width)

def add_p(doc, text="", bold=False, italic=False, size=12, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return p

def add_h1(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_H1
    return p

def add_h2(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_H2
    return p

def add_h3(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_H3
    return p

def add_callout(doc, text, title="KEY ANALYTICAL INSIGHT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_TITLE
    
    r_text = p.add_run(text)
    r_text.font.name = FONT_NAME
    r_text.font.size = Pt(9.5)
    r_text.font.italic = True
    r_text.font.color.rgb = COLOR_TEXT
    
    # Empty space after callout
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_figure_with_caption(doc, fig_filename, fig_num, caption_title, width_inches=6.0):
    fig_path = os.path.join(FIG_DIR, fig_filename)
    if not os.path.exists(fig_path):
        print(f"Warning: Figure {fig_filename} not found!")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(fig_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    
    r_label = p_cap.add_run(f"Figure {fig_num}: ")
    r_label.font.name = FONT_NAME
    r_label.font.size = Pt(9.5)
    r_label.font.bold = True
    r_label.font.color.rgb = COLOR_TITLE
    
    r_text = p_cap.add_run(caption_title)
    r_text.font.name = FONT_NAME
    r_text.font.size = Pt(9.5)
    r_text.font.italic = True
    r_text.font.color.rgb = COLOR_TEXT

def add_table_caption(doc, table_num, caption_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    r_label = p.add_run(f"Table {table_num}: ")
    r_label.font.name = FONT_NAME
    r_label.font.size = Pt(9.5)
    r_label.font.bold = True
    r_label.font.color.rgb = COLOR_TITLE
    
    r_text = p.add_run(caption_title)
    r_text.font.name = FONT_NAME
    r_text.font.size = Pt(9.5)
    r_text.font.italic = True
    r_text.font.color.rgb = COLOR_TEXT

def add_dynamic_page_number_to_footer(section):
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_p.paragraph_format.space_before = Pt(0)
    f_p.paragraph_format.space_after = Pt(0)
    
    # Add dynamic page number field in Word XML
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    t_val = parse_xml(f'<w:t {nsdecls("w")}>1</w:t>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    
    r = f_p.add_run()
    r.font.name = FONT_NAME
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_MUTED
    
    f_p._p.append(fldChar1)
    f_p._p.append(instrText)
    f_p._p.append(fldChar2)
    f_p._p.append(t_val)
    f_p._p.append(fldChar3)

print("Helper functions compiled successfully.")
