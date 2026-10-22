"""
build_master_report.py
Master orchestrator to build the complete Technical Report Word Document:
Traffic_Accident_Analysis_and_Severity_Risk_Modeling_Technical_Report.docx
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

import report_sections_front_and_lit as front_lit
import report_sections_intro_and_method as intro_method
import report_sections_results_part1 as res_part1
import report_sections_results_part2 as res_part2
import report_sections_conclusion_ref as conc_ref
import build_full_report as helpers

OUTPUT_DOCX = os.path.join(BASE_DIR, "Traffic_Accident_Analysis_and_Severity_Risk_Modeling_Technical_Report.docx")

def build_complete_report():
    print("================================================================================")
    print("BUILDING COMPREHENSIVE TECHNICAL REPORT (DOCX)")
    print("================================================================================")
    
    doc = docx.Document()
    
    # Configure 1.25" left margin for academic binding, 1.0" top/bottom/right
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.25)
        s.right_margin = Inches(1.0)
        helpers.add_dynamic_page_number_to_footer(s)
        
    print("1. Building Cover Page...")
    front_lit.build_cover_page(doc, helpers.add_p)
    
    print("2. Building Abstract & Keywords...")
    front_lit.build_abstract(doc, helpers.add_h1, helpers.add_p, helpers.add_callout)
    
    print("3. Building Table of Contents...")
    front_lit.build_table_of_contents(doc, helpers.add_h1, helpers.add_table_data)
    
    print("4. Building List of Figures...")
    front_lit.build_list_of_figures(doc, helpers.add_h1, helpers.add_p)
    
    print("5. Building List of Tables...")
    front_lit.build_list_of_tables(doc, helpers.add_h1, helpers.add_p)
    
    print("6. Building Section 1: Introduction...")
    intro_method.build_section_1_introduction(doc, helpers.add_h1, helpers.add_h2, helpers.add_p, helpers.add_callout)
    
    print("7. Building Section 2: Literature Review...")
    intro_method.build_section_2_literature_review(doc, helpers.add_h1, helpers.add_h2, helpers.add_p, helpers.add_callout)
    
    print("8. Building Section 3: Research Questions...")
    intro_method.build_section_3_research_questions(doc, helpers.add_h1, helpers.add_h2, helpers.add_p, helpers.add_callout)
    
    print("9. Building Section 4: Methodology & System Architecture...")
    intro_method.build_section_4_methodology(doc, helpers.add_h1, helpers.add_h2, helpers.add_h3, helpers.add_p, helpers.add_callout)
    
    print("10. Building Section 5 Part 1: Results 5.1 to 5.4 (Figures 1-11, Tables 1-11)...")
    res_part1.build_section_5_part1(doc, helpers.add_h1, helpers.add_h2, helpers.add_h3, helpers.add_p, helpers.add_callout, 
                                   helpers.add_figure_with_caption, helpers.add_table_caption, helpers.add_table_data)
    
    print("11. Building Section 5 Part 2: Results 5.5 to 5.8 (Figures 12-22, Tables 12-22)...")
    res_part2.build_section_5_part2(doc, helpers.add_h2, helpers.add_h3, helpers.add_p, helpers.add_callout, 
                                   helpers.add_figure_with_caption, helpers.add_table_caption, helpers.add_table_data)
    
    print("12. Building Section 6: Conclusions & Recommendations...")
    conc_ref.build_section_6_conclusions(doc, helpers.add_h1, helpers.add_h2, helpers.add_p, helpers.add_callout)
    
    print("13. Building References...")
    conc_ref.build_references(doc, helpers.add_h1, helpers.add_p)
    
    print("14. Building Appendix...")
    conc_ref.build_appendix(doc, helpers.add_h1, helpers.add_h2, helpers.add_p)
    
    print(f"15. Saving complete document to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    
    file_size_mb = os.path.getsize(OUTPUT_DOCX) / (1024 * 1024)
    print("================================================================================")
    print(f"SUCCESS! Technical Report successfully created: {OUTPUT_DOCX}")
    print(f"File Size: {file_size_mb:.2f} MB")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print("================================================================================")

if __name__ == "__main__":
    build_complete_report()
